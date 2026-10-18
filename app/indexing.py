from langchain_community.document_loaders import Docx2txtLoader
from langchain_community.document_loaders import PyPDFLoader
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_community.embeddings import OllamaEmbeddings
from langchain_community.vectorstores import Chroma
import os
import shutil
from docx import Document
from langchain.schema import Document as LangchainDocument


def custom_arabic_text_splitter_by_heading1(docx_path: str):
    doc = Document(docx_path)
    sections = []
    current_section = ""

    for para in doc.paragraphs:
        if para.style.name == "Heading 1":
            if current_section:
                sections.append(current_section)
                current_section = ""

            current_section += para.text + "\n"
        else:
            current_section += para.text + "\n"

    if current_section:
        sections.append(current_section)

    return sections


def clear_old_store(persist_directory):
    if os.path.exists(persist_directory):
        shutil.rmtree(persist_directory)


def create_english_vector_store():
    pdf_path = "data/raw/en.pdf"
    clear_old_store("data/vector_store/en") 
    loader = PyPDFLoader(pdf_path)
    docs = loader.load()

    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=200
    )

    splits = text_splitter.split_documents(docs)

    embeddings = OllamaEmbeddings(
        model="nomic-embed-text"
    )

    vectorstore_en = Chroma.from_documents(
        documents=splits,
        embedding=embeddings,
        persist_directory="data/vector_store/en"
    )

    vectorstore_en.persist()

    return vectorstore_en


def create_arabic_vector_store():
    docx_path = "data/raw/ar.docx"
    clear_old_store("data/vector_store/ar")
    sections = custom_arabic_text_splitter_by_heading1(docx_path)

    splits = [
        LangchainDocument(page_content=section)
        for section in sections
    ]

    embeddings = HuggingFaceEmbeddings(
        model_name="sentence-transformers/paraphrase-multilingual-mpnet-base-v2"
    )

    vectorstore_ar = Chroma.from_documents(
        documents=splits,
        embedding=embeddings,
        persist_directory="data/vector_store/ar"
    )

    vectorstore_ar.persist()

    return vectorstore_ar


def create_vector_stores():
    print("Creating English vector database...")
    create_english_vector_store()

    print("English vector database created successfully.")

    print("Creating Arabic vector database...")
    create_arabic_vector_store()

    print("Arabic vector database created successfully.")

    print("All vector databases created successfully.")


if __name__ == "__main__":
    create_vector_stores()
